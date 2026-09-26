"""Extract the supplied 2022 workbook; never evaluate instructions in its cells."""
import json, re, sys
from pathlib import Path
import openpyxl
book = openpyxl.load_workbook(sys.argv[1], data_only=True)
expansions = {'Inspect. J.3':'Inspecteur J.3','Compt. Contr. J.3':'Comptable Contrôleur J.3','Ten. de Livre 3':'Teneur de Livre 3','Commis ADM 3':'Commis Administratif 3','Dactilog. 2':'Dactylographe 2','Opér. Inf. 3':'Opérateur informatique 3','Opér. Inf. 4':'Opérateur informatique 4','Ménagère 2 et Autres':'Agent d’entretien 2','Ménagère 1':'Agent d’entretien 1','Auxiliaire Infirmière':'Auxiliaire infirmier','Infirmière de ligne':'Infirmier de ligne','Infirmière spécialiste':'Infirmier spécialiste','Chef Service':'Chef de Service','Assistant Administratif (1)':'Assistant Administratif 1'}
pairs = {'Premier':'Première','Assistant':'Assistante','Administrateur':'Administratrice','Chargé':'Chargée','Chef':'Cheffe','Conseiller':'Conseillère','Coordonnateur':'Coordonnatrice','Directeur':'Directrice','Départemental':'Départementale','Doyen':'Doyenne','Ingénieur':'Ingénieure','Consultant':'Consultante','Vérificateur':'Vérificatrice','Délégué':'Déléguée','Adjoint':'Adjointe','Administratif':'Administrative','Inspecteur':'Inspectrice','Programmeur':'Programmeuse','Professionnel':'Professionnelle','Contrôleur':'Contrôleuse','Financier':'Financière','Agent':'Agente','Douanier':'Douanière','douanier':'douanière','Technicien':'Technicienne','Teneur':'Teneuse','Opérateur':'Opératrice','Intendant':'Intendante','Chauffeur':'Chauffeuse','Mécanicien':'Mécanicienne','Messager':'Messagère','infirmier':'infirmière','Infirmier':'Infirmière','Médical':'Médicale','Pharmacien':'Pharmacienne','Magasinier':'Magasinière','Régisseur':'Régisseuse'}
def feminine(s):
 return re.sub(r'[^\W\d_]+',lambda m:pairs.get(m[0],m[0]),s)
entries={}
for row in range(6,54):
 cat,sub,titles,_,_,salary,_,page,note=[book.worksheets[0].cell(row,c).value for c in range(1,10)]
 if cat=='Personnel policier': continue
 category=sub if cat=='Autres personnels administratifs et techniques' else cat
 for raw in re.split(r',\s*|\s+/\s+',titles):
  # These catch-all descriptions are not job titles and must not approve arbitrary salaries.
  if raw.startswith(('Autres ','Salaire Minimum')): continue
  label=expansions.get(raw,raw)
  key=(category,label)
  if key not in entries:
   entries[key]={'id':f'grid-2022-{row}-{len(entries)+1}','masculine':label,'feminine':feminine(label),'category':category,'salaries':[],'aliases':[],'source':'Grille salariale Haïti — mai 2022','sourceRows':[],'notes':'','active':True,'version':1}
  e=entries[key]
  if salary not in e['salaries']: e['salaries'].append(salary)
  if raw!=label and raw not in e['aliases']: e['aliases'].append(raw)
  e['sourceRows'].append(row)
  if note and ('incertaine' in note or 'confirmer' in note): e['notes']=note
# Undated annotations are retained but not used as approved salaries until activated.
for row,male,cat in [(6,'Magasinier','Personnel de soutien'),(7,'Régisseur de pharmacie','Personnel médical'),(8,'Archiviste','Personnel administratif'),(9,'Massothérapeute','Personnel médical'),(10,'Manutentionnaire','Personnel de soutien')]:
 raw,salary,_,note=[book.worksheets[1].cell(row,c).value for c in range(1,5)]
 entries[('annotation',male)]={'id':f'grid-note-{row}','masculine':male,'feminine':feminine(male),'category':cat,'salaries':[salary],'aliases':[raw] if raw!=male else [],'source':'Annotations manuscrites — sans date','sourceRows':[row],'notes':note or 'Montant manuscrit non daté, à valider.','active':False,'version':1}
output=Path('src/features/salary-grid/salaryGridSeed.json')
output.write_text(json.dumps(list(entries.values()),ensure_ascii=False,indent=2)+'\n')
print(f'{len(entries)} titres, {sum(e["active"] for e in entries.values())} actifs')
